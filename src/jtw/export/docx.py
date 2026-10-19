"""Word export: the final English, a Japanese/English side-by-side draft, or the source alone.

Segments joined by ``merge_with_next`` become one paragraph. English is typeset (curly
quotes, ellipses) and ``**bold**`` / ``*italic*`` become Word runs. The page header carries the
title, author, part label and translator from ``project.toml``.
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from enum import StrEnum
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

from jtw.config import Project
from jtw.export.typography import runs, smarten
from jtw.model.segment import Segment, merge_chains
from jtw.model.store import LoadedPart


class ExportKind(StrEnum):
    FINAL = "final"
    SIDE_BY_SIDE = "side-by-side"
    SOURCE = "source"


_SUFFIX = {ExportKind.FINAL: "", ExportKind.SIDE_BY_SIDE: ".JE_Draft", ExportKind.SOURCE: ".JA"}
_HEADER_NOTE = {
    ExportKind.FINAL: "",
    ExportKind.SIDE_BY_SIDE: " (J-E draft)",
    ExportKind.SOURCE: " (source)",
}


def export_filename(project: Project, part_id: str, kind: ExportKind) -> str:
    """``<Title>.<PartId>[.JE_Draft|.JA].docx`` with the title made filename-safe."""
    title = re.sub(r"[^\w-]+", "_", project.config.work.title).strip("_") or project.slug
    return f"{title}.{part_id}{_SUFFIX[kind]}.docx"


def joined_final(chain: Sequence[Segment]) -> str:
    return " ".join(s.final.strip() for s in chain if s.final.strip())


def joined_source(chain: Sequence[Segment]) -> str:
    return "".join(s.source.strip() for s in chain)


def build_document(project: Project, part: LoadedPart, kind: ExportKind) -> Any:
    document = Document()
    _header(document, project, part.part.label + _HEADER_NOTE[kind], kind)
    for chain in merge_chains(list(part.segments.values())):
        english = joined_final(chain)
        if kind is ExportKind.SOURCE:
            _japanese(document, joined_source(chain), space_after=12)
            continue
        if not english:
            continue
        if kind is ExportKind.SIDE_BY_SIDE:
            _japanese(document, joined_source(chain), space_after=4)
        _english(document, english)
    return document


def export_part(
    project: Project, part: LoadedPart, kind: ExportKind, out_dir: Path | None = None
) -> Path:
    directory = out_dir if out_dir is not None else project.paths.exports
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / export_filename(project, part.part.id, kind)
    build_document(project, part, kind).save(str(path))
    return path


def _header(document: Any, project: Project, label: str, kind: ExportKind) -> None:
    work = project.config.work
    title = work.title_source if kind is ExportKind.SOURCE and work.title_source else work.title
    right_bottom = f"Trans. {work.translator}" if work.translator else ""
    header = document.sections[0].header
    header.paragraphs[0].text = ""
    _header_line(header.paragraphs[0], title, work.author, bold_left=True)
    _header_line(header.add_paragraph(), label, right_bottom, bold_left=False)


def _header_line(paragraph: Any, left: str, right: str, *, bold_left: bool) -> None:
    stops = paragraph.paragraph_format.tab_stops
    stops.add_tab_stop(Inches(3), WD_TAB_ALIGNMENT.CENTER)
    stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT)
    run = paragraph.add_run(left)
    run.bold = bold_left
    run.font.size = Pt(12)
    if right:
        paragraph.add_run("\t\t")
        paragraph.add_run(right).font.size = Pt(12)
    paragraph.paragraph_format.space_after = Pt(0)


def _english(document: Any, text: str) -> None:
    paragraph = document.add_paragraph()
    for run in runs(smarten(text)):
        added = paragraph.add_run(run.text)
        added.bold = run.bold or None
        added.italic = run.italic or None
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(12)
    fmt.first_line_indent = Inches(0.5)


def _japanese(document: Any, text: str, *, space_after: int) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(space_after)
