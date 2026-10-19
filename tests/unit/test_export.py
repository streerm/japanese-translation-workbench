from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from jtw.config import Project, load_project
from jtw.export.clipboard import chains_for_range, copy_range, render, to_system_clipboard
from jtw.export.docx import ExportKind, export_filename, export_part
from jtw.export.stats import format_table, project_stats
from jtw.export.typography import Run, runs, smarten
from jtw.model import store
from jtw.model.segment import Segment, merge_chains
from jtw.paths import ProjectPaths


@pytest.fixture
def project(worked_paths: ProjectPaths) -> Project:
    return load_project("mini")


def _part(project: Project, part_id: str = "C01P01") -> store.LoadedPart:
    return store.load_part(project.paths, part_id, project.config.columns)


def test_smarten() -> None:
    assert smarten('"Wait..." she said. "Don\'t."') == "“Wait…” she said. “Don’t.”"
    assert smarten("the 'night' (\"x\")") == "the \u2018night\u2019 (\u201cx\u201d)"


def test_runs() -> None:
    assert list(runs("a **b** *c* ***d*** e")) == [
        Run("a "),
        Run("b", bold=True),
        Run(" "),
        Run("c", italic=True),
        Run(" "),
        Run("d", bold=True, italic=True),
        Run(" e"),
    ]
    assert list(runs("2 * 3 = 6")) == [Run("2 * 3 = 6")]


def test_merge_chains() -> None:
    segs = [Segment(f"S00{i}", merge_with_next=i in (2, 3)) for i in range(1, 6)]
    assert [[s.key for s in c] for c in merge_chains(segs)] == [
        ["S001"],
        ["S002", "S003", "S004"],
        ["S005"],
    ]
    last_flagged = [Segment("S001", merge_with_next=True)]
    assert len(merge_chains(last_flagged)) == 1


def _paragraphs(path: Path) -> list[object]:
    return [p for p in Document(str(path)).paragraphs if p.text.strip()]


def test_final_export(project: Project, tmp_path: Path) -> None:
    path = export_part(project, _part(project), ExportKind.FINAL, tmp_path)
    assert path.name == "Mini_Story.C01P01.docx"
    paras = _paragraphs(path)
    texts = [p.text for p in paras]  # type: ignore[attr-defined]
    assert texts == [
        "I woke at seven…",
        "Outside, the rain fell quietly. “Morning,” my sister said.",
        "I made toast and tea.",
        "I decided to walk to the station.",
    ]
    italic = [r.text for r in paras[1].runs if r.italic]  # type: ignore[attr-defined]
    bold = [r.text for r in paras[2].runs if r.bold]  # type: ignore[attr-defined]
    assert italic == ["quietly"] and bold == ["toast"]
    header = Document(str(path)).sections[0].header.paragraphs
    assert header[0].text == "Mini Story\t\tTest Author"
    assert header[1].text == "Chapter 1 · Part 1\t\tTrans. Test Translator"


def test_side_by_side_and_source_exports(project: Project) -> None:
    sbs = export_part(project, _part(project), ExportKind.SIDE_BY_SIDE)
    assert sbs.parent == project.paths.exports and sbs.name.endswith(".JE_Draft.docx")
    texts = [p.text for p in _paragraphs(sbs)]  # type: ignore[attr-defined]
    assert len(texts) == 8
    assert texts[2] == "窓の外では雨が静かに降っていた。「おはよう」と姉が言った。"
    assert texts[3].startswith("Outside, the rain fell quietly.")

    src = export_part(project, _part(project), ExportKind.SOURCE)
    texts = [p.text for p in _paragraphs(src)]  # type: ignore[attr-defined]
    assert len(texts) == 4  # S002+S003 and S005+S006 chains
    assert texts[-1] == "テーブルの上に新聞があった。今日は駅まで歩くことにした。"
    assert Document(str(src)).sections[0].header.paragraphs[0].text.startswith("小さな話")
    assert export_filename(project, "P01", ExportKind.SOURCE) == "Mini_Story.P01.JA.docx"


def test_untranslated_part_exports_empty_body(project: Project) -> None:
    path = export_part(project, _part(project, "C01P02"), ExportKind.FINAL)
    assert _paragraphs(path) == []


def test_clipboard_render_drops_empty_lines() -> None:
    template = "Source: {source}\nComment: {comment}\nFixed line"
    assert render(template, {"source": "源", "comment": " "}) == "Source: 源\nFixed line"


def test_copy_range_full_expands_merge_chains(project: Project) -> None:
    part = _part(project)
    assert [[s.key for s in c] for c in chains_for_range(part, "S003", "S003")] == [
        ["S002", "S003"]
    ]
    text = copy_range(project, part, "S003")
    assert text == (
        "Source: 窓の外では雨が静かに降っていた。「おはよう」と姉が言った。\n"
        "Model A: Alpha S002. Alpha S003.\n"
        "Model B: \n"
        'Current Draft: Outside, the rain fell *quietly*. "Morning," my sister said.\n'
        "Comment: [S003] Sister or older sister?"
    )


def test_copy_range_minimal_multiple_blocks(project: Project) -> None:
    text = copy_range(project, _part(project), "S004", "S001", mode="minimal")
    blocks = text.split("\n\n")
    assert len(blocks) == 3
    assert blocks[0] == "Source: 朝の七時に目が覚めた。\nCurrent Draft: I woke at seven..."
    assert "Comment" not in text


def test_system_clipboard_failure_is_reported(monkeypatch: pytest.MonkeyPatch) -> None:
    import pyperclip

    def broken(text: str) -> None:
        raise pyperclip.PyperclipException("no clipboard")

    monkeypatch.setattr(pyperclip, "copy", broken)
    assert to_system_clipboard("x") is False
    monkeypatch.setattr(pyperclip, "copy", lambda text: None)
    assert to_system_clipboard("x") is True


def test_stats(project: Project) -> None:
    stats = project_stats(project)
    first, second = stats.parts
    assert (first.segments, second.segments) == (6, 6)
    assert first.chars == 11 + 16 + 13 + 16 + 14 + 14
    assert first.finalized_segments == 2
    assert first.finalized_chars == 11 + 13
    assert first.translated == {"alpha": 6, "beta": 0}
    assert second.translated == {"alpha": 0, "beta": 0} and second.finalized_pct == 0
    assert stats.chars == first.chars + second.chars
    assert round(stats.finalized_pct, 1) == round(100 * 24 / stats.chars, 1)
    table = format_table(stats).splitlines()
    assert table[0].split()[:4] == ["part", "segs", "chars", "final"]
    assert table[1].startswith("C01P01")
    assert table[-1].startswith("total")
    data = stats.to_dict()
    assert data["totals"]["segments"] == 12
    assert data["parts"][0]["translated"] == {"alpha": 6, "beta": 0}
